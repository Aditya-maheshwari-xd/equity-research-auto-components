"""Build the Endurance Technologies DCF workbook.

Every forecast, valuation and sensitivity cell is a live Excel formula;
inputs (blue) are documented with their source next to them.

    python dcf/build_dcf.py            -> dcf/Endurance_DCF.xlsx
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("Endurance_DCF.xlsx")

# --------------------------------------------------------------------------
# Source data (INR crore unless noted). StockAnalysis.com (S&P Global Market
# Intelligence) reports INR millions; divided by 10 here.
# --------------------------------------------------------------------------
SRC_SA_IS = "StockAnalysis.com income statement (S&P Global data), stockanalysis.com/quote/nse/ENDURANCE/financials/income-statement/"
SRC_SA_BS = "StockAnalysis.com balance sheet, stockanalysis.com/quote/nse/ENDURANCE/financials/balance-sheet/"
SRC_SA_CF = "StockAnalysis.com cash flow, stockanalysis.com/quote/nse/ENDURANCE/financials/cash-flow-statement/"

YEARS_H = ["FY22", "FY23", "FY24", "FY25", "FY26"]
HIST = {  # INR millions from StockAnalysis
    "Revenue": ([75491, 88040, 102409, 115608, 145959], SRC_SA_IS),
    "EBITDA": ([9414, 10189, 13071, 15317, 19311], SRC_SA_IS),
    "D&A": ([3585, 4042, 4531, 5193, 6997], SRC_SA_IS),
    "EBIT": ([5829, 6147, 8540, 10124, 12314], SRC_SA_IS),
    "Net income": ([4607, 4796, 6805, 8364, 9517], SRC_SA_IS),
    "Operating cash flow": ([7416, 8620, 10571, 15317, 18507], SRC_SA_CF),
    "Capex": ([5290, 6363, 8348, 10534, 12958], SRC_SA_CF),
    "Acquisitions": ([59.29, 1206, 69.43, 203.4, 3576], SRC_SA_CF),
    "Receivables": ([11367, 13477, 15197, 16933, 21453], SRC_SA_BS),
    "Inventory": ([7287, 8653, 9153, 9844, 12472], SRC_SA_BS),
    "Payables": ([12413, 14257, 16045, 17950, 23693], SRC_SA_BS),
    "Cash & ST investments": ([8882, 9340, 12697, 18068, 18894], SRC_SA_BS),
    "Total debt (incl. leases)": ([4296, 5148, 7653, 9437, 13267], SRC_SA_BS),
    "Shareholders' equity": ([39200, 44121, 49774, 57174, 68408], SRC_SA_BS),
}
EPS = ([32.75, 34.09, 48.38, 59.46, 67.66], SRC_SA_IS)

EQ = "Shareholders' equity"
YEARS_F = [f"FY{y}E" for y in range(27, 37)]
N = len(YEARS_F)
SCEN_GROWTH = {
    "Bear": [0.12, 0.09, 0.09, 0.08, 0.08, 0.07, 0.07, 0.06, 0.06, 0.06],
    "Base": [0.17, 0.14, 0.13, 0.12, 0.11, 0.10, 0.10, 0.09, 0.09, 0.08],
    "Bull": [0.21, 0.17, 0.16, 0.15, 0.14, 0.13, 0.12, 0.11, 0.10, 0.10],
}
SCEN_MARGIN = {
    "Bear": [0.124, 0.124, 0.125, 0.125, 0.125, 0.125, 0.125, 0.125, 0.125, 0.125],
    "Base": [0.130, 0.133, 0.136, 0.138, 0.138, 0.138, 0.138, 0.138, 0.138, 0.138],
    "Bull": [0.134, 0.139, 0.143, 0.146, 0.148, 0.148, 0.148, 0.148, 0.148, 0.148],
}
CAPEX_PCT = [0.080, 0.075, 0.070, 0.065, 0.062, 0.060, 0.058, 0.056, 0.055, 0.055]

# --------------------------------------------------------------------------
# Styles
# --------------------------------------------------------------------------
F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
BOLD_GREEN = Font(name=F, size=10, bold=True, color="008000")
TITLE = Font(name=F, size=14, bold=True, color="1F3A5F")
SUB = Font(name=F, size=9, italic=True, color="666666")
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
OUT_FILL = PatternFill("solid", fgColor="E8EEF6")
THIN = Side(style="thin", color="999999")
TOP = Border(top=THIN)
CR = '#,##0;(#,##0);"-"'
CR1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
RS = '"₹"#,##0;("₹"#,##0);"-"'
RS2 = '"₹"#,##0.00'


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, note=None, bold=False, border=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=F, size=10, bold=bold or font.bold, color=font.color, italic=font.italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if note:
        c.comment = Comment(note, "Model")
    return c


def header_row(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = HDR
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center" if i else "left")


def widths(ws, first=34, rest=11, n=14, extra=None):
    ws.column_dimensions["A"].width = first
    for i in range(2, n + 2):
        ws.column_dimensions[L(i)].width = rest
    for col, w in (extra or {}).items():
        ws.column_dimensions[col].width = w


wb = Workbook()

# ==========================================================================
# Historicals
# ==========================================================================
hs = wb.active
hs.title = "Historicals"
widths(hs, extra={"G": 70})
put(hs, "A1", "Endurance Technologies: consolidated historicals", TITLE)
put(hs, "A2", "INR crore, fiscal years ending 31 March. Blue = reported figures (source in column G).", SUB)
header_row(hs, 4, ["INR crore"] + YEARS_H + ["Source"])
H = {}  # label -> row
r = 5
for label, (vals, src) in HIST.items():
    put(hs, f"A{r}", label)
    for j, v in enumerate(vals):
        put(hs, f"{L(2 + j)}{r}", round(v / 10, 1), BLUE, CR)
    put(hs, f"G{r}", src, SUB)
    H[label] = r
    r += 1
put(hs, f"A{r}", "EPS (INR)")
for j, v in enumerate(EPS[0]):
    put(hs, f"{L(2 + j)}{r}", v, BLUE, RS2)
put(hs, f"G{r}", EPS[1], SUB)
H["EPS"] = r
r += 2

put(hs, f"A{r}", "Ratios (formulas)", BOLD)
r += 1
ratio_rows = [
    ("Revenue growth", lambda c, p: f"=IF({p}{H['Revenue']}=0,0,{c}{H['Revenue']}/{p}{H['Revenue']}-1)", True),
    ("EBITDA margin", lambda c, p: f"={c}{H['EBITDA']}/{c}{H['Revenue']}", False),
    ("D&A % revenue", lambda c, p: f"={c}{H['D&A']}/{c}{H['Revenue']}", False),
    ("Capex % revenue", lambda c, p: f"={c}{H['Capex']}/{c}{H['Revenue']}", False),
    ("Net working capital (Rec + Inv - Pay)",
     lambda c, p: f"={c}{H['Receivables']}+{c}{H['Inventory']}-{c}{H['Payables']}", False),
    ("NWC % revenue", None, False),
    ("Effective tax proxy: 1 - NI/EBIT", None, False),
    ("ROE (NI / closing equity)", lambda c, p: f"={c}{H['Net income']}/{c}{H[EQ]}", False),
    ("Free cash flow (OCF - capex)", lambda c, p: f"={c}{H['Operating cash flow']}-{c}{H['Capex']}", False),
    ("Net cash / (debt)", lambda c, p: f"={c}{H['Cash & ST investments']}-{c}{H['Total debt (incl. leases)']}", False),
]
for label, fn, skip_first in ratio_rows:
    put(hs, f"A{r}", label)
    H[label] = r
    for j in range(5):
        col, prev = L(2 + j), L(1 + j)
        if label == "NWC % revenue":
            f = f"={col}{H['Net working capital (Rec + Inv - Pay)']}/{col}{H['Revenue']}"
        elif label.startswith("Effective tax"):
            f = f"=1-{col}{H['Net income']}/{col}{H['EBIT']}"
        elif skip_first and j == 0:
            continue
        else:
            f = fn(col, prev)
        is_pct = label not in ("Net working capital (Rec + Inv - Pay)", "Free cash flow (OCF - capex)",
                               "Net cash / (debt)")
        put(hs, f"{col}{r}", f, BLACK, PCT if is_pct else CR)
    r += 1
put(hs, f"A{r + 1}", "Note: EBITDA here excludes other income (S&P definition). Company-reported FY26 EBITDA "
                     "including other income was ₹2,090 Cr (14.2% of total income), per the FY26 press release "
                     "(NSE filing, 14 May 2026).", SUB)
put(hs, f"A{r + 2}", "Note: 'NI / EBIT' mixes interest, other income and tax, so it is only a sanity check; "
                     "the model uses the statutory 25.17% rate (Section 115BAA).", SUB)

# ==========================================================================
# WACC
# ==========================================================================
w = wb.create_sheet("WACC")
widths(w, first=40, rest=13, n=3, extra={"C": 90})
put(w, "A1", "Discount rate (WACC)", TITLE)
put(w, "A2", "INR, nominal. Blue = input with source; black = formula.", SUB)
header_row(w, 4, ["Item", "Value", "Source / reasoning"])
W = {}
rows = [
    ("India 10Y G-sec yield", 0.0718, PCT, "Trading Economics, 28 Sep 2026: India 10Y near a two-year high at ~7.18%."),
    ("India sovereign default spread", 0.0187, PCT, "Damodaran country risk table (updated 5 Jan 2026): India Baa3, adjusted default spread 1.87%."),
    ("INR risk-free rate", "=B5-B6", PCT, "Govt yield less default spread, so country risk isn't counted twice (it is already in the ERP below)."),
    ("Levered beta", 0.87, '0.00', "sharpely.in (price date 16 Oct 2026). Below 1: auto parts are cyclical, but ENDU is net cash. Flexed in Sensitivity."),
    ("Equity risk premium (India, total)", 0.0708, PCT, "Damodaran (5 Jan 2026): mature-market ERP 4.23% + India CRP 2.85%."),
    ("Cost of equity (CAPM)", "=B7+B8*B9", PCT, "Rf + beta x ERP"),
    ("Pre-tax cost of debt", 0.075, PCT, "Assumption. FY26 interest / avg debt is ~5% (cheap Euro loans + leases); 7.5% reflects marginal INR borrowing for an AA+ issuer (ICRA AA+ Stable, Sep 2025)."),
    ("Tax rate", "=Assumptions!C27", PCT, "Linked to Assumptions"),
    ("After-tax cost of debt", "=B11*(1-B12)", PCT, ""),
    ("Share price (INR)", 2406.40, RS2, "NSE close 16 Oct 2026 (Dhan / 5paisa / Value Research)."),
    ("Shares outstanding (crore)", "=B16/B14", '0.000', "Implied from market cap / price."),
    ("Market cap (INR Cr)", 33849, CR, "Dhan / 5paisa, 16 Oct 2026: ₹33,849 Cr."),
    ("Total debt incl. leases (INR Cr)", "=Historicals!F" + str(H["Total debt (incl. leases)"]), CR, "FY26 balance sheet"),
    ("Equity weight", "=B16/(B16+B17)", PCT, "Market values"),
    ("Debt weight", "=B17/(B16+B17)", PCT, ""),
    ("WACC", "=B18*B10+B19*B13", PCT, ""),
]
for i, (label, val, fmt, note) in enumerate(rows):
    rr = 5 + i
    W[label] = rr
    is_input = not (isinstance(val, str) and val.startswith("="))
    font = BLUE if is_input else (GREEN if "!" in str(val) else BLACK)
    put(w, f"A{rr}", label, BOLD if label in ("WACC", "Cost of equity (CAPM)") else BLACK)
    put(w, f"B{rr}", val, font, fmt, fill=KEY_FILL if label == "WACC" else None)
    put(w, f"C{rr}", note, SUB)
assert W["India 10Y G-sec yield"] == 5 and W["WACC"] == 20

# ==========================================================================
# Assumptions
# ==========================================================================
a = wb.create_sheet("Assumptions", 0)
widths(a, first=36, rest=10, n=13, extra={"M": 60})
put(a, "A1", "Forecast assumptions", TITLE)
put(a, "A2", "Blue cells are inputs. Change the scenario in C4 (1 = Bear, 2 = Base, 3 = Bull) and the whole model updates.", SUB)
put(a, "A4", "Scenario selected", BOLD)
put(a, "C4", 2, BLUE, "0", KEY_FILL, note="1 = Bear, 2 = Base, 3 = Bull")
put(a, "D4", '=CHOOSE(C4,"Bear","Base","Bull")', BOLD)
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", showErrorMessage=True,
                    error="Enter 1, 2 or 3")
a.add_data_validation(dv)
dv.add("C4")
header_row(a, 6, ["Driver", "FY26A"] + YEARS_F + ["Rationale"])

A = {}


def scen_block(start, title, data, live_label, hist_formula, rationale):
    put(a, f"A{start}", title, BOLD)
    for k, (name, vals) in enumerate(data.items()):
        rr = start + 1 + k
        put(a, f"A{rr}", f"  {name}")
        for j, v in enumerate(vals):
            put(a, f"{L(3 + j)}{rr}", v, BLUE, PCT)
    live = start + 4
    put(a, f"A{live}", live_label, BOLD)
    put(a, f"B{live}", hist_formula, GREEN, PCT)
    for j in range(N):
        col = L(3 + j)
        put(a, f"{col}{live}", f"=CHOOSE($C$4,{col}{start + 1},{col}{start + 2},{col}{start + 3})", BOLD, PCT,
            fill=OUT_FILL)
    put(a, f"M{start}", rationale, SUB)
    return live


A["growth"] = scen_block(
    7, "Revenue growth", SCEN_GROWTH, "Revenue growth (live)",
    f"=Historicals!F{H['Revenue growth']}",
    "Base: Q1 FY27 sales +29.6% (Stoferle, strong 2W/PV volumes); FY27E 17% assumes H2 normalises on a high base. "
    "Then fades to 8% by FY36: a 10-year CAGR of ~11%, in line with the company's own 10-year sales CAGR (11%, Screener). Bull adds ABS/alloy-wheel/4W wins; Bear = 2W slowdown + weak Europe.")
A["margin"] = scen_block(
    13, "EBITDA margin (excl. other income)", SCEN_MARGIN, "EBITDA margin (live)",
    f"=Historicals!F{H['EBITDA margin']}",
    "FY26 13.2%, TTM Jun-26 12.9% (input costs). Base recovers to 13.8% from higher-margin ABS, alloy wheels and "
    "Europe EV/hybrid programmes. Bear stays ~12.5%.")

put(a, "A20", "Other drivers (all scenarios)", BOLD)
other = [
    (21, "D&A % revenue", [0.047] * N, f"=Historicals!F{H['D&A % revenue']}",
     "FY26 4.8%; held near that as new plants depreciate."),
    (22, "Capex % revenue", CAPEX_PCT, f"=Historicals!F{H['Capex % revenue']}",
     "FY26 8.9% (expansion year). Mgmt: FY27 capex similar to FY26. Falls to 5.5% as plants fill; stays above D&A to fund growth."),
    (23, "NWC % revenue", [0.07] * N, f"=Historicals!F{H['NWC % revenue']}",
     "FY26 7.0% (Rec + Inv - Pay); 5-yr range 7-9%. Held at 7%."),
]
for rr, label, vals, hist, why in other:
    put(a, f"A{rr}", label)
    put(a, f"B{rr}", hist, GREEN, PCT)
    for j, v in enumerate(vals):
        put(a, f"{L(3 + j)}{rr}", v, BLUE, PCT)
    put(a, f"M{rr}", why, SUB)
    A[label] = rr

put(a, "A26", "Valuation inputs", BOLD)
inputs = [
    (27, "Tax rate", 0.2517, PCT, "Statutory rate under Sec 115BAA (22% + surcharge + cess)."),
    (28, "Terminal growth", 0.05, PCT, "Nominal INR; below India's long-run nominal GDP (~10%) — assumes share of growth slows."),
    (29, "Valuation date (base) ", "31-Mar-2026", None, "Cash flows discounted from FY26 year end (mid-year convention)."),
    (30, "Days from 31-Mar-26 to 16-Oct-26", 199, "0", "Equity value is rolled forward at cost of equity to the price date."),
    (31, "FY26 final dividend paid (INR/share)", 11.5, RS2, "FY26 press release: ₹11.50/share recommended; paid before price date, so deducted."),
    (32, "Stoferle: price paid for 8% (EUR m)", 6.24, '0.00', "Q1 FY27 investor presentation: EUR 6.24m for an additional 8% (stake now 68%)."),
    (33, "Stoferle: remaining stake to buy", 0.32, PCT, "Remaining 32% to be acquired over four years (investor presentation)."),
    (34, "EUR/INR", "=1146*10/104.3", "0.00", "Implied from Q1 FY27 Europe revenue: ₹1,146 Cr (= ₹11,460m) / EUR 104.3m."),
    (35, "Stoferle future payment (INR Cr) — treated like debt", "=B32/0.08*B33*B34/10", CR,
     "Implied value of 100% (6.24/8%) x 32% x FX. EUR m x INR/EUR / 10 = INR crore."),
]
for rr, label, val, fmt, why in inputs:
    put(a, f"A{rr}", label)
    formula = isinstance(val, str) and val.startswith("=")
    put(a, f"B{rr}" if rr >= 29 else f"C{rr}", val, BLACK if formula else BLUE, fmt)
    put(a, f"M{rr}", why, SUB)
# tax and g in column C for consistency with WACC links
A["tax"] = "Assumptions!$C$27"
A["g"] = "Assumptions!$C$28"

# ==========================================================================
# DCF
# ==========================================================================
d = wb.create_sheet("DCF", 1)
widths(d, first=38, rest=11, n=12)
put(d, "A1", "Endurance Technologies: discounted cash flow (FCFF)", TITLE)
put(d, "A2", '="INR crore. Scenario: "&Assumptions!D4&". Mid-year discounting from 31-Mar-2026."', SUB)
header_row(d, 4, ["INR crore", "FY26A"] + YEARS_F)
D = {}
labels = ["Revenue", "  growth", "EBITDA", "  margin", "D&A", "EBIT", "Tax on EBIT", "NOPAT", "+ D&A", "- Capex",
          "- Increase in NWC", "Free cash flow to firm (FCFF)", "", "Discount period (years, mid-year)",
          "Discount factor", "PV of FCFF"]
for i, lab in enumerate(labels):
    D[lab] = 5 + i
    put(d, f"A{5 + i}", lab, BOLD if lab in ("Revenue", "EBITDA", "Free cash flow to firm (FCFF)", "PV of FCFF") else BLACK)

hf = "Historicals!F"
put(d, f"B{D['Revenue']}", f"={hf}{H['Revenue']}", GREEN, CR)
put(d, f"B{D['  growth']}", f"={hf}{H['Revenue growth']}", GREEN, PCT)
put(d, f"B{D['EBITDA']}", f"={hf}{H['EBITDA']}", GREEN, CR)
put(d, f"B{D['  margin']}", f"=B{D['EBITDA']}/B{D['Revenue']}", BLACK, PCT)
put(d, f"B{D['D&A']}", f"={hf}{H['D&A']}", GREEN, CR)
put(d, f"B{D['EBIT']}", f"=B{D['EBITDA']}-B{D['D&A']}", BLACK, CR)
put(d, f"B{D['- Capex']}", f"=-{hf}{H['Capex']}", GREEN, CR)

for j in range(N):
    c, p = L(3 + j), L(2 + j)
    ac = c  # same column letters on Assumptions (C..L)
    put(d, f"{c}{D['  growth']}", f"=Assumptions!{ac}{A['growth']}", GREEN, PCT)
    put(d, f"{c}{D['Revenue']}", f"={p}{D['Revenue']}*(1+{c}{D['  growth']})", BLACK, CR)
    put(d, f"{c}{D['  margin']}", f"=Assumptions!{ac}{A['margin']}", GREEN, PCT)
    put(d, f"{c}{D['EBITDA']}", f"={c}{D['Revenue']}*{c}{D['  margin']}", BLACK, CR)
    put(d, f"{c}{D['D&A']}", f"={c}{D['Revenue']}*Assumptions!{ac}{A['D&A % revenue']}", BLACK, CR)
    put(d, f"{c}{D['EBIT']}", f"={c}{D['EBITDA']}-{c}{D['D&A']}", BLACK, CR)
    put(d, f"{c}{D['Tax on EBIT']}", f"=-{c}{D['EBIT']}*{A['tax']}", BLACK, CR)
    put(d, f"{c}{D['NOPAT']}", f"={c}{D['EBIT']}+{c}{D['Tax on EBIT']}", BLACK, CR)
    put(d, f"{c}{D['+ D&A']}", f"={c}{D['D&A']}", BLACK, CR)
    put(d, f"{c}{D['- Capex']}", f"=-{c}{D['Revenue']}*Assumptions!{ac}{A['Capex % revenue']}", BLACK, CR)
    put(d, f"{c}{D['- Increase in NWC']}",
        f"=-({c}{D['Revenue']}-{p}{D['Revenue']})*Assumptions!{ac}{A['NWC % revenue']}", BLACK, CR)
    put(d, f"{c}{D['Free cash flow to firm (FCFF)']}",
        f"={c}{D['NOPAT']}+{c}{D['+ D&A']}+{c}{D['- Capex']}+{c}{D['- Increase in NWC']}", BOLD, CR, border=TOP)
    put(d, f"{c}{D['Discount period (years, mid-year)']}", j + 0.5, BLACK, '0.0')
    put(d, f"{c}{D['Discount factor']}",
        f"=1/(1+WACC!$B$20)^{c}{D['Discount period (years, mid-year)']}", BLACK, '0.000')
    put(d, f"{c}{D['PV of FCFF']}", f"={c}{D['Free cash flow to firm (FCFF)']}*{c}{D['Discount factor']}", BOLD, CR)
d[f"A{D['Discount period (years, mid-year)']}"].comment = Comment(
    "Hard-coded 0.5, 1.5, ... : cash assumed to arrive mid-year.", "Model")

FC = f"C{D['Free cash flow to firm (FCFF)']}:L{D['Free cash flow to firm (FCFF)']}"
PER = f"C{D['Discount period (years, mid-year)']}:L{D['Discount period (years, mid-year)']}"
last_fcff = f"L{D['Free cash flow to firm (FCFF)']}"

v0 = D["PV of FCFF"] + 2
put(d, f"A{v0}", "Valuation", BOLD)
V = {}
val_rows = [
    ("Sum of PV of FCFF (FY27-FY36)", f"=SUM(C{D['PV of FCFF']}:L{D['PV of FCFF']})", CR),
    ("WACC", "=WACC!B20", PCT),
    ("Terminal growth", f"={A['g']}", PCT),
    ("Terminal value (Gordon growth, at FY36)", f"={last_fcff}*(1+B{{g}})/(B{{wacc}}-B{{g}})", CR),
    ("PV of terminal value", f"=B{{tv}}/(1+B{{wacc}})^{N}", CR),
    ("Enterprise value (31-Mar-26)", "=B{sum}+B{pvtv}", CR),
    ("Terminal value % of EV", "=B{pvtv}/B{ev}", PCT),
    ("+ Cash & short-term investments", f"={hf}{H['Cash & ST investments']}", CR),
    ("- Total debt incl. leases", f"=-{hf}{H['Total debt (incl. leases)']}", CR),
    ("- Stoferle remaining stake (minority)", "=-Assumptions!B35", CR),
    ("Equity value (31-Mar-26)", "=SUM(B{ev},B{cash},B{debt},B{mi})", CR),
    ("Roll forward to 16-Oct-26 at cost of equity", "=B{eq0}*(1+WACC!B10)^(Assumptions!B30/365)", CR),
    ("- FY26 dividend paid", "=-Assumptions!B31*WACC!B15", CR),
    ("Equity value (16-Oct-26)", "=B{roll}+B{div}", CR),
    ("Shares outstanding (crore)", "=WACC!B15", '0.000'),
    ("Intrinsic value per share (INR)", "=B{eq1}/B{sh}", RS),
    ("Current share price (16-Oct-26)", "=WACC!B14", RS),
    ("Upside / (downside)", "=B{ps}/B{px}-1", PCT),
]
keys = ["sum", "wacc", "g", "tv", "pvtv", "ev", "tvpct", "cash", "debt", "mi", "eq0", "roll", "div", "eq1", "sh",
        "ps", "px", "up"]
rowmap = {k: v0 + 1 + i for i, k in enumerate(keys)}
for i, ((label, f, fmt), k) in enumerate(zip(val_rows, keys)):
    rr = rowmap[k]
    V[k] = rr
    formula = f.format(**rowmap)
    green = formula.count("!") and not any(x in formula for x in ("+", "*", "/", "SUM(")) or k in ("cash", "debt")
    key = k in ("ps", "up")
    put(d, f"A{rr}", label, BOLD if key or k in ("ev", "eq1") else BLACK)
    put(d, f"B{rr}", formula, GREEN if green else BLACK, fmt, fill=KEY_FILL if key else None,
        bold=key or k in ("ev", "eq1"))

m0 = V["up"] + 2
put(d, f"A{m0}", "Implied multiples at intrinsic value", BOLD)
put(d, f"A{m0 + 1}", "EV / FY27E EBITDA")
put(d, f"B{m0 + 1}", f"=(B{V['eq1']}-B{V['cash']}-B{V['debt']}-B{V['mi']})/C{D['EBITDA']}", BLACK, MULT)
put(d, f"A{m0 + 2}", "Implied terminal EV / EBITDA (FY36)")
put(d, f"B{m0 + 2}", f"=B{V['tv']}/L{D['EBITDA']}", BLACK, MULT)
put(d, f"A{m0 + 3}", "Market EV / FY27E EBITDA (current price)")
put(d, f"B{m0 + 3}", f"=(WACC!B16-B{V['cash']}-B{V['debt']}-B{V['mi']})/C{D['EBITDA']}", BLACK, MULT)
put(d, f"A{m0 + 5}", "Checks", BOLD)
put(d, f"A{m0 + 6}", "WACC > terminal growth")
put(d, f"B{m0 + 6}", f'=IF(B{V["wacc"]}>B{V["g"]},"OK","ERROR")', BOLD)
put(d, f"A{m0 + 7}", "Terminal capex >= D&A (growth must be funded)")
put(d, f"B{m0 + 7}", f'=IF(Assumptions!L{A["Capex % revenue"]}>=Assumptions!L{A["D&A % revenue"]},"OK","CHECK")', BOLD)
d.freeze_panes = "B5"

# ==========================================================================
# Sensitivity
# ==========================================================================
s = wb.create_sheet("Sensitivity", 2)
widths(s, first=30, rest=11, n=8)
put(s, "A1", "Sensitivity of intrinsic value per share (INR)", TITLE)
put(s, "A2", "Every cell re-runs the valuation with the FCFF stream from the DCF sheet (current scenario).", SUB)

ev_core = (f"(SUMPRODUCT(DCF!{FC}/(1+{{w}})^DCF!{PER})"
           f"+DCF!{last_fcff}*(1+{{g}})/({{w}}-{{g}})/(1+{{w}})^{N})")
bridge = (f"((({{ev}})+DCF!$B${V['cash']}+DCF!$B${V['debt']}+DCF!$B${V['mi']})"
          f"*(1+WACC!$B$10)^(Assumptions!$B$30/365)+DCF!$B${V['div']})/DCF!$B${V['sh']}")

put(s, "A4", "WACC (down) x terminal growth (across)", BOLD)
gs = [0.04, 0.045, 0.05, 0.055, 0.06]
ws_ = [0.0975, 0.1025, 0.1075, 0.1125, 0.1175, 0.1225, 0.1275]
for j, g in enumerate(gs):
    put(s, f"{L(2 + j)}5", g, BLUE, PCT, fill=OUT_FILL)
for i, wv in enumerate(ws_):
    rr = 6 + i
    put(s, f"A{rr}", wv, BLUE, PCT, fill=OUT_FILL)
    for j in range(len(gs)):
        col = L(2 + j)
        ev = ev_core.format(w=f"$A{rr}", g=f"{col}$5")
        put(s, f"{col}{rr}", "=" + bridge.format(ev=ev), BLACK, RS)
put(s, "A14", "Current price", BOLD)
put(s, "B14", "=WACC!B14", GREEN, RS)
put(s, "A15", "Live WACC in model", BOLD)
put(s, "B15", "=WACC!B20", GREEN, PCT)

put(s, "A17", "Cross-check: exit multiple instead of Gordon growth", BOLD)
put(s, "A18", "Terminal EV / FY36 EBITDA", BOLD)
mults = [10, 12, 14, 16, 18]
for j, m in enumerate(mults):
    put(s, f"{L(2 + j)}18", m, BLUE, MULT, fill=OUT_FILL)
put(s, "A19", "Value per share (INR)")
for j in range(len(mults)):
    col = L(2 + j)
    ev = (f"SUMPRODUCT(DCF!{FC}/(1+WACC!$B$20)^DCF!{PER})"
          f"+{col}$18*DCF!$L${D['EBITDA']}/(1+WACC!$B$20)^{N}")
    put(s, f"{col}19", "=" + bridge.format(ev=ev), BLACK, RS)
put(s, "A20", "Implied terminal growth")
for j in range(len(mults)):
    col = L(2 + j)
    # TV = FCFF*(1+g)/(w-g) -> g = (TV*w - FCFF)/(TV + FCFF)
    tv = f"{col}$18*DCF!$L${D['EBITDA']}"
    put(s, f"{col}20", f"=({tv}*WACC!$B$20-DCF!{last_fcff})/({tv}+DCF!{last_fcff})", BLACK, PCT)

put(s, "A22", "Beta (down) x equity risk premium (across): WACC", BOLD)
betas = [0.7, 0.87, 1.0, 1.2]
erps = [0.06, 0.0708, 0.08]
for j, e in enumerate(erps):
    put(s, f"{L(2 + j)}23", e, BLUE, PCT, fill=OUT_FILL)
for i, b in enumerate(betas):
    rr = 24 + i
    put(s, f"A{rr}", b, BLUE, '0.00', fill=OUT_FILL)
    for j in range(len(erps)):
        col = L(2 + j)
        put(s, f"{col}{rr}", f"=WACC!$B$18*(WACC!$B$7+$A{rr}*{col}$23)+WACC!$B$19*WACC!$B$13", BLACK, PCT)

# ==========================================================================
# Comps
# ==========================================================================
cp = wb.create_sheet("Comps", 3)
widths(cp, first=30, rest=12, n=8, extra={"H": 46})
put(cp, "A1", "Listed Indian auto-component peers", TITLE)
put(cp, "A2", "Market data as reported by Value Research / 5paisa on the date shown (INR). Blue = sourced input.", SUB)
header_row(cp, 4, ["Company", "Price", "Mkt cap (Cr)", "P/E (TTM)", "EV/EBITDA", "ROE", "As of", "Note"])
peers = [
    ("Endurance Technologies", 2406.40, 33829, 34.9, 15.6, 0.149, "16-Oct-26", "Value Research; ROE from Screener"),
    ("Uno Minda", 1108.00, 64110, 53.3, 28.7, 0.205, "16-Oct-26", "Lighting, switches, alloy wheels, EV parts"),
    ("Sona BLW Precision", 800.00, 49859, 71.6, 38.7, 0.111, "16-Oct-26", "EV driveline; premium for EV exposure"),
    ("Gabriel India", 1303.60, 25075, 84.2, 51.6, 0.198, "05-Oct-26", "Closest product peer (suspension)"),
    ("Minda Corporation", 649.85, 15537, 31.0, None, None, "16-Oct-26", "5paisa; EV/EBITDA, ROE not shown"),
    ("Bharat Forge", 1833.00, 89129, 126.2, 30.0, 0.116, "05-Oct-26", "Forgings + defence; P/E distorted, excluded from median"),
]
for i, row in enumerate(peers):
    rr = 5 + i
    for j, v in enumerate(row):
        fmt = [None, RS2, CR, MULT, MULT, PCT, None, None][j]
        font = BOLD if (i == 0 and j == 0) else (BLUE if j in (1, 2, 3, 4, 5) and v is not None else BLACK)
        if v is None:
            v = "n/a"
        put(cp, f"{L(1 + j)}{rr}", v, font, fmt)
put(cp, "A12", "Peer median (excl. Endurance, Bharat Forge)", BOLD)
put(cp, "D12", "=MEDIAN(D6:D9)", BOLD, MULT)
put(cp, "E12", "=MEDIAN(E6:E8)", BOLD, MULT)
put(cp, "F12", "=MEDIAN(F6:F8)", BOLD, PCT)
put(cp, "A13", "Endurance discount to peer median", BOLD)
put(cp, "D13", "=D5/D12-1", BOLD, PCT)
put(cp, "E13", "=E5/E12-1", BOLD, PCT)
put(cp, "A15", "Implied value per share at peer median P/E x Endurance TTM EPS (₹68.95, StockAnalysis TTM Jun-26)")
put(cp, "D15", "=D12*68.95", BLACK, RS)
put(cp, "A16", "Implied value at Endurance's own TTM P/E applied to FY27E EPS proxy is in the pitch deck.", SUB)
put(cp, "A18", "Sources: valueresearchonline.com stock pages (Uno Minda, Sona BLW, Gabriel India, Bharat Forge, "
                "Endurance key ratios); 5paisa.com (Minda Corp); screener.in (Endurance ROE). Retrieved 18 Oct 2026.", SUB)

# ==========================================================================
# Cover
# ==========================================================================
cv = wb.create_sheet("Cover", 0)
cv.column_dimensions["A"].width = 44
cv.column_dimensions["B"].width = 18
cv.column_dimensions["C"].width = 70
put(cv, "A1", "Endurance Technologies Ltd (NSE: ENDURANCE)", TITLE)
put(cv, "A2", "DCF valuation model, prepared 18 Oct 2026. Student research for learning purposes, not investment advice.", SUB)
summary = [
    ("Scenario", "=Assumptions!D4", None),
    ("Intrinsic value per share", f"=DCF!B{V['ps']}", RS),
    ("Share price (NSE close, 16-Oct-26)", "=WACC!B14", RS),
    ("Upside / (downside)", f"=DCF!B{V['up']}", PCT),
    ("WACC", "=WACC!B20", PCT),
    ("Terminal growth", f"={A['g']}", PCT),
    ("Terminal value % of EV", f"=DCF!B{V['tvpct']}", PCT),
    ("Enterprise value (INR Cr)", f"=DCF!B{V['ev']}", CR),
]
for i, (k, f, fmt) in enumerate(summary):
    put(cv, f"A{4 + i}", k, BOLD)
    put(cv, f"B{4 + i}", f, GREEN, fmt, fill=KEY_FILL if "Intrinsic" in k or "Upside" in k else None)
put(cv, "A13", "How to use", BOLD)
howto = [
    "Assumptions!C4 switches scenario: 1 = Bear, 2 = Base, 3 = Bull.",
    "Blue = input (each has a source or rationale beside it). Black = formula. Green = link to another sheet.",
    "Yellow = key outputs and the scenario switch.",
    "Sensitivity tab re-runs the valuation for different WACC, terminal growth and exit multiples.",
    "Historicals: StockAnalysis.com (S&P Global). Company filings on NSE for FY26 and Q1 FY27 results.",
]
for i, t in enumerate(howto):
    put(cv, f"A{14 + i}", t)
put(cv, "A20", "Sheets", BOLD)
for i, (nm, desc) in enumerate([("Assumptions", "scenario switch and every forecast driver"),
                                 ("DCF", "10-year FCFF forecast, terminal value, equity bridge"),
                                 ("Sensitivity", "WACC x g, exit multiple, beta x ERP"),
                                 ("Comps", "listed peer multiples"), ("WACC", "cost of capital build-up"),
                                 ("Historicals", "FY22-FY26 financials and ratios")]):
    put(cv, f"A{21 + i}", nm)
    put(cv, f"B{21 + i}", desc)

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT)
print(f"Wrote {OUT}")
print("rows:", {"V": V, "D": D})
